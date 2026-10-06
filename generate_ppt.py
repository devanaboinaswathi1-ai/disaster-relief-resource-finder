"""
Script to generate a 15-slide PowerPoint presentation for Disaster Relief Finder project
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Title Slide
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Header
    header = slide1.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    header_frame = header.text_frame
    header_frame.text = "DISASTER RELIEF RESOURCE FINDER"
    header_para = header_frame.paragraphs[0]
    header_para.font.size = Pt(32)
    header_para.font.bold = True
    header_para.font.color.rgb = RGBColor(45, 212, 191)
    header_para.alignment = PP_ALIGN.CENTER
    
    # Title at center
    title = slide1.shapes.add_textbox(Inches(2), Inches(3), Inches(6), Inches(1))
    title_frame = title.text_frame
    title_frame.text = "Web-Based Disaster Relief Resource Locator"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(24)
    title_para.font.bold = True
    title_para.alignment = PP_ALIGN.CENTER
    
    # Guide name at bottom left
    guide = slide1.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(4), Inches(0.8))
    guide_frame = guide.text_frame
    guide_frame.text = "Guide: [Guide Name]\nDesignation: [Professor/Assistant Professor]"
    guide_para = guide_frame.paragraphs[0]
    guide_para.font.size = Pt(12)
    
    # Student details at top right
    student = slide1.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(3), Inches(1))
    student_frame = student.text_frame
    student_frame.text = "Project Done By:\n[Student Name]\nRoll No: [Roll Number]"
    student_para = student_frame.paragraphs[0]
    student_para.font.size = Pt(12)
    student_para.alignment = PP_ALIGN.RIGHT
    
    # Slide 2: Table of Contents
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    title2 = slide2.shapes.title
    title2.text = "TABLE OF CONTENTS"
    
    content2 = slide2.placeholders[1].text_frame
    content2.clear()
    points = [
        "Introduction to Project",
        "Objectives",
        "Advantages",
        "Abstract & Keywords",
        "Existing Systems & Drawbacks",
        "Proposed System & Merits",
        "Hardware & Software Requirements",
        "System Architecture"
    ]
    for point in points:
        p = content2.add_paragraph()
        p.text = point
        p.level = 0
        p.font.size = Pt(16)
    
    # Slide 3: Abstract
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    title3 = slide3.shapes.title
    title3.text = "ABSTRACT"
    
    content3 = slide3.placeholders[1].text_frame
    content3.clear()
    p = content3.add_paragraph()
    p.text = "The Disaster Relief Resource Finder is a web-based GIS application that helps citizens locate essential relief resources during natural disasters and emergencies."
    p.font.size = Pt(14)
    
    p = content3.add_paragraph()
    p.text = "Main Functionality:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    features = [
        "Interactive map-based resource locator using Leaflet.js",
        "Real-time location tracking with GPS integration",
        "Filter relief centers by type (shelter, food, medical, water)",
        "Radius-based search (1-100 km) for nearby resources",
        "Admin dashboard for CRUD operations on relief centers",
        "Responsive design for mobile and desktop access"
    ]
    for feature in features:
        p = content3.add_paragraph()
        p.text = feature
        p.level = 1
        p.font.size = Pt(12)
    
    p = content3.add_paragraph()
    p.text = "Keywords: Disaster Management, GIS, Relief Centers, Emergency Response, Web Mapping, PHP, MySQL, Leaflet.js"
    p.font.size = Pt(11)
    p.font.italic = True
    
    # Slide 4: Existing Systems & Drawbacks
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    title4 = slide4.shapes.title
    title4.text = "EXISTING SYSTEMS & DRAWBACKS"
    
    content4 = slide4.placeholders[1].text_frame
    content4.clear()
    
    p = content4.add_paragraph()
    p.text = "Existing Systems:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    existing = [
        "Google Crisis Map - Global disaster tracking",
        "NDRF Mobile App - National disaster response",
        "State Government Portals - Static information pages",
        "Social Media Updates - Unverified crowdsourced data"
    ]
    for item in existing:
        p = content4.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    p = content4.add_paragraph()
    p.text = "Drawbacks:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    drawbacks = [
        "Limited local coverage for specific regions",
        "No real-time admin updates for relief centers",
        "Complex interfaces not user-friendly during emergencies",
        "Lack of radius-based proximity search",
        "No integrated CRUD operations for administrators",
        "Poor mobile responsiveness"
    ]
    for item in drawbacks:
        p = content4.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    # Slide 5: Proposed System & Merits
    slide5 = prs.slides.add_slide(prs.slide_layouts[1])
    title5 = slide5.shapes.title
    title5.text = "PROPOSED SYSTEM & MERITS"
    
    content5 = slide5.placeholders[1].text_frame
    content5.clear()
    
    p = content5.add_paragraph()
    p.text = "Proposed System:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    p = content5.add_paragraph()
    p.text = "A comprehensive web-based platform with interactive mapping, real-time updates, and location-based search for disaster relief resources across Andhra Pradesh and major Indian cities."
    p.font.size = Pt(12)
    
    p = content5.add_paragraph()
    p.text = "Advantages Over Existing Systems:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    merits = [
        "Localized focus with 100+ relief centers in AP",
        "Real-time admin dashboard for instant updates",
        "Simple, intuitive interface for emergency use",
        "Customizable radius search (1-100 km)",
        "Complete CRUD operations for data management",
        "Fully responsive mobile-first design",
        "Open-source and easily deployable",
        "Offline-capable with service worker support (future)"
    ]
    for item in merits:
        p = content5.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    # Slide 6: Hardware & Software Requirements
    slide6 = prs.slides.add_slide(prs.slide_layouts[1])
    title6 = slide6.shapes.title
    title6.text = "HARDWARE & SOFTWARE REQUIREMENTS"
    
    content6 = slide6.placeholders[1].text_frame
    content6.clear()
    
    p = content6.add_paragraph()
    p.text = "Hardware Requirements:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    hardware = [
        "Processor: Intel Core i3 or higher",
        "RAM: 4GB minimum, 8GB recommended",
        "Storage: 500MB for application files",
        "Network: Internet connection for map tiles"
    ]
    for item in hardware:
        p = content6.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    p = content6.add_paragraph()
    p.text = "Software Requirements:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    software = [
        "Server: XAMPP/WAMP/LAMP (Apache 2.4+)",
        "Backend: PHP 7.4 or higher",
        "Database: MySQL 5.7+ / MariaDB",
        "Frontend: HTML5, CSS3, JavaScript ES6+",
        "Mapping Library: Leaflet.js 1.9.4",
        "Browser: Chrome, Firefox, Safari, Edge (latest)",
        "Development Tools: VS Code, phpMyAdmin"
    ]
    for item in software:
        p = content6.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    # Slide 7: Functional & Non-Functional Requirements
    slide7 = prs.slides.add_slide(prs.slide_layouts[1])
    title7 = slide7.shapes.title
    title7.text = "FUNCTIONAL & NON-FUNCTIONAL REQUIREMENTS"
    
    content7 = slide7.placeholders[1].text_frame
    content7.clear()
    
    p = content7.add_paragraph()
    p.text = "Functional Requirements:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    functional = [
        "User location detection via GPS/manual selection",
        "Search relief centers by type and radius",
        "Display centers on interactive map with markers",
        "Admin authentication and session management",
        "CRUD operations: Add, Edit, Delete relief centers",
        "Distance calculation using Haversine formula",
        "Filter and sort results by proximity"
    ]
    for item in functional:
        p = content7.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    p = content7.add_paragraph()
    p.text = "Non-Functional Requirements:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    nonfunctional = [
        "Performance: Page load < 3 seconds",
        "Scalability: Support 1000+ concurrent users",
        "Security: SQL injection prevention, session security",
        "Usability: Intuitive UI, accessible during stress",
        "Reliability: 99.5% uptime, error handling",
        "Compatibility: Cross-browser and mobile responsive"
    ]
    for item in nonfunctional:
        p = content7.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(12)
    
    # Slide 8: System Architecture
    slide8 = prs.slides.add_slide(prs.slide_layouts[1])
    title8 = slide8.shapes.title
    title8.text = "SYSTEM ARCHITECTURE"
    
    content8 = slide8.placeholders[1].text_frame
    content8.clear()
    
    p = content8.add_paragraph()
    p.text = "Three-Tier Architecture:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    p = content8.add_paragraph()
    p.text = "1. Presentation Layer (Frontend):"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content8.add_paragraph()
    p.text = "HTML5, CSS3, JavaScript, Leaflet.js for interactive UI"
    p.level = 1
    p.font.size = Pt(11)
    
    p = content8.add_paragraph()
    p.text = "2. Application Layer (Backend):"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content8.add_paragraph()
    p.text = "PHP scripts for API endpoints (get_centers.php, add_center.php, update_center.php, delete_center.php)"
    p.level = 1
    p.font.size = Pt(11)
    
    p = content8.add_paragraph()
    p.text = "3. Data Layer (Database):"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content8.add_paragraph()
    p.text = "MySQL database with 'centers' table storing relief center information"
    p.level = 1
    p.font.size = Pt(11)
    
    p = content8.add_paragraph()
    p.text = "Components: User Browser → Apache Server → PHP Backend → MySQL DB → OpenStreetMap Tiles"
    p.font.size = Pt(11)
    p.font.italic = True
    
    # Slide 9: System Design Diagrams
    slide9 = prs.slides.add_slide(prs.slide_layouts[1])
    title9 = slide9.shapes.title
    title9.text = "SYSTEM DESIGN DIAGRAMS"
    
    content9 = slide9.placeholders[1].text_frame
    content9.clear()
    
    p = content9.add_paragraph()
    p.text = "Use Case Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "Actors: Public User, Admin | Use Cases: Search Centers, View Map, Filter by Type, Add/Edit/Delete Centers, Login"
    p.level = 1
    p.font.size = Pt(10)
    
    p = content9.add_paragraph()
    p.text = "Class Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "Classes: Center (id, name, type, lat, lng, address, phone), User, Admin, Map"
    p.level = 1
    p.font.size = Pt(10)
    
    p = content9.add_paragraph()
    p.text = "Sequence Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "User → UI → API → Database → Response flow for search and CRUD operations"
    p.level = 1
    p.font.size = Pt(10)
    
    p = content9.add_paragraph()
    p.text = "Activity Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "Flow: Start → Select Location → Choose Type → Set Radius → Search → Display Results → End"
    p.level = 1
    p.font.size = Pt(10)
    
    p = content9.add_paragraph()
    p.text = "Deployment Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "Client Browser ↔ Apache Web Server ↔ PHP Runtime ↔ MySQL Server"
    p.level = 1
    p.font.size = Pt(10)
    
    p = content9.add_paragraph()
    p.text = "Component Diagram:"
    p.font.size = Pt(13)
    p.font.bold = True
    p = content9.add_paragraph()
    p.text = "Components: Map Module, Search Module, Admin Module, Database Module, API Module"
    p.level = 1
    p.font.size = Pt(10)
    
    # Slide 10: Module Description
    slide10 = prs.slides.add_slide(prs.slide_layouts[1])
    title10 = slide10.shapes.title
    title10.text = "MODULE DESCRIPTION"
    
    content10 = slide10.placeholders[1].text_frame
    content10.clear()
    
    modules = [
        ("User Interface Module", "Interactive map display, search controls, filter options, location selector, results list panel"),
        ("Location Module", "GPS geolocation, manual city selection, coordinate conversion, distance calculation"),
        ("Search & Filter Module", "Type-based filtering, radius-based search, proximity sorting, real-time results"),
        ("Mapping Module", "Leaflet.js integration, marker placement, popup information, map tile loading"),
        ("Admin Authentication Module", "Login system, session management, access control, logout functionality"),
        ("CRUD Operations Module", "Add new centers, edit existing data, delete centers, validate inputs"),
        ("Database Module", "MySQL connection, prepared statements, data retrieval, transaction management"),
        ("API Module", "RESTful endpoints, JSON responses, error handling, CORS configuration")
    ]
    
    for title, desc in modules:
        p = content10.add_paragraph()
        p.text = f"{title}:"
        p.font.size = Pt(12)
        p.font.bold = True
        p = content10.add_paragraph()
        p.text = desc
        p.level = 1
        p.font.size = Pt(10)
    
    # Slide 11: Implementation
    slide11 = prs.slides.add_slide(prs.slide_layouts[1])
    title11 = slide11.shapes.title
    title11.text = "IMPLEMENTATION"
    
    content11 = slide11.placeholders[1].text_frame
    content11.clear()
    
    p = content11.add_paragraph()
    p.text = "Frontend Implementation:"
    p.font.size = Pt(13)
    p.font.bold = True
    
    frontend = [
        "HTML5 semantic structure with responsive meta tags",
        "CSS3 with custom properties for theming",
        "JavaScript ES6+ with async/await for API calls",
        "Leaflet.js for map initialization and marker management",
        "Event listeners for user interactions"
    ]
    for item in frontend:
        p = content11.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(11)
    
    p = content11.add_paragraph()
    p.text = "Backend Implementation:"
    p.font.size = Pt(13)
    p.font.bold = True
    
    backend = [
        "PHP 7.4+ with MySQLi extension",
        "Prepared statements to prevent SQL injection",
        "JSON encoding for API responses",
        "Session management for admin authentication",
        "Haversine formula for distance calculation",
        "Error handling with HTTP status codes"
    ]
    for item in backend:
        p = content11.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(11)
    
    # Slide 12: Testing
    slide12 = prs.slides.add_slide(prs.slide_layouts[1])
    title12 = slide12.shapes.title
    title12.text = "TESTING"
    
    content12 = slide12.placeholders[1].text_frame
    content12.clear()
    
    p = content12.add_paragraph()
    p.text = "Test Cases Executed:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    tests = [
        "TC01: Database connection test - PASSED",
        "TC02: User location detection via GPS - PASSED",
        "TC03: Manual city selection from dropdown - PASSED",
        "TC04: Filter by resource type (shelter/food/medical/water) - PASSED",
        "TC05: Radius-based search (1-100 km) - PASSED",
        "TC06: Map marker display and popup - PASSED",
        "TC07: Distance calculation accuracy - PASSED",
        "TC08: Admin login authentication - PASSED",
        "TC09: Add new relief center - PASSED",
        "TC10: Edit existing center - PASSED",
        "TC11: Delete center with confirmation - PASSED",
        "TC12: API error handling - PASSED",
        "TC13: Mobile responsiveness - PASSED",
        "TC14: Cross-browser compatibility - PASSED",
        "TC15: SQL injection prevention - PASSED"
    ]
    for test in tests:
        p = content12.add_paragraph()
        p.text = test
        p.level = 1
        p.font.size = Pt(11)
    
    # Slide 13: Conclusion & Future Enhancements
    slide13 = prs.slides.add_slide(prs.slide_layouts[1])
    title13 = slide13.shapes.title
    title13.text = "CONCLUSION & FUTURE ENHANCEMENTS"
    
    content13 = slide13.placeholders[1].text_frame
    content13.clear()
    
    p = content13.add_paragraph()
    p.text = "Conclusion:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    p = content13.add_paragraph()
    p.text = "Successfully developed a web-based disaster relief resource finder with interactive mapping, real-time updates, and location-based search. The system provides quick access to 100+ relief centers across Andhra Pradesh, helping citizens during emergencies."
    p.font.size = Pt(12)
    
    p = content13.add_paragraph()
    p.text = "Future Enhancements:"
    p.font.size = Pt(14)
    p.font.bold = True
    
    future = [
        "Real-time updates using WebSockets",
        "Mobile app development (React Native/Flutter)",
        "SMS/Email notification system",
        "Multi-language support (Telugu, Hindi, Tamil)",
        "Offline mode with service workers",
        "Advanced analytics dashboard",
        "Integration with emergency services APIs",
        "Crowdsourced updates and verification",
        "AI-based resource allocation prediction"
    ]
    for item in future:
        p = content13.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(11)
    
    # Slide 14: Queries
    slide14 = prs.slides.add_slide(prs.slide_layouts[1])
    title14 = slide14.shapes.title
    title14.text = "QUERIES"
    
    content14 = slide14.placeholders[1].text_frame
    content14.clear()
    
    p = content14.add_paragraph()
    p.text = "Questions & Answers"
    p.font.size = Pt(24)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    p = content14.add_paragraph()
    p.text = " "
    
    p = content14.add_paragraph()
    p.text = "Thank you for your attention!"
    p.font.size = Pt(20)
    p.alignment = PP_ALIGN.CENTER
    
    p = content14.add_paragraph()
    p.text = " "
    
    p = content14.add_paragraph()
    p.text = "Feel free to ask any questions about:"
    p.font.size = Pt(14)
    
    queries = [
        "System architecture and design",
        "Implementation details",
        "Technology stack choices",
        "Testing methodology",
        "Future enhancements",
        "Deployment process"
    ]
    for item in queries:
        p = content14.add_paragraph()
        p.text = item
        p.level = 1
        p.font.size = Pt(13)
    
    # Slide 15: References
    slide15 = prs.slides.add_slide(prs.slide_layouts[1])
    title15 = slide15.shapes.title
    title15.text = "REFERENCES"
    
    content15 = slide15.placeholders[1].text_frame
    content15.clear()
    
    references = [
        "Leaflet.js Documentation - https://leafletjs.com/",
        "OpenStreetMap - https://www.openstreetmap.org/",
        "PHP Official Documentation - https://www.php.net/docs.php",
        "MySQL Documentation - https://dev.mysql.com/doc/",
        "National Disaster Management Authority (NDMA) - https://ndma.gov.in/",
        "State Disaster Management Authority, Andhra Pradesh",
        "Web GIS Applications for Disaster Management - Research Papers",
        "HTML5 Geolocation API - MDN Web Docs",
        "RESTful API Design Best Practices",
        "Responsive Web Design Principles - W3C Standards"
    ]
    
    for ref in references:
        p = content15.add_paragraph()
        p.text = ref
        p.font.size = Pt(12)
        p.level = 0
    
    # Save presentation
    prs.save('Disaster_Relief_Finder_Presentation.pptx')
    print("✓ PowerPoint presentation created successfully!")
    print("✓ File saved as: Disaster_Relief_Finder_Presentation.pptx")
    print("✓ Total slides: 15")

if __name__ == "__main__":
    try:
        create_presentation()
    except ImportError:
        print("Error: python-pptx library not found!")
        print("Please install it using: pip install python-pptx")
    except Exception as e:
        print(f"Error creating presentation: {e}")
